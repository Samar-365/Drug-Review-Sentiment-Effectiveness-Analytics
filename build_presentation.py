import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Constants
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)    # Slate 50
    CARD_BG = RGBColor(255, 255, 255)     # White
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    PRIMARY = RGBColor(30, 41, 59)        # Slate 800
    ACCENT_BLUE = RGBColor(14, 165, 233)  # Sky 500
    ACCENT_TEAL = RGBColor(16, 185, 129)  # Emerald 500
    ACCENT_AMBER = RGBColor(245, 158, 11) # Amber 500
    ACCENT_PURPLE = RGBColor(139, 92, 246)# Violet 500
    ACCENT_RED = RGBColor(239, 68, 68)    # Red 500
    TEXT_MAIN = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="DRUG REVIEW SENTIMENT ANALYTICS"):
        # Top accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT_BLUE
        bar.line.color.rgb = ACCENT_BLUE

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, left_stripe=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
            
        if left_stripe:
            stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.12), Inches(height))
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = left_stripe
            stripe.line.fill.background()
        return card

    def add_arrow(slide, left, top, width=0.3, height=0.25, color=ACCENT_BLUE):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background()
        return arrow

    def add_down_arrow(slide, left, top, width=0.25, height=0.22, color=ACCENT_BLUE):
        arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background()
        return arrow

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, BG_DARK)

    tbar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    tbar.fill.solid()
    tbar.fill.fore_color.rgb = ACCENT_BLUE
    tbar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.3), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "HONOURS PROJECT: DATA SCIENCE & HEALTHCARE"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p1 = tf.add_paragraph()
    p1.text = "Drug Review Sentiment & Effectiveness Analytics"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Patient Feedback Intelligence & Clinical Sentiment Platform"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(8)

    members = [
        ("Samar", "Project Lead & ML Architect", "Sample Dataset, Architecture & Releases", ACCENT_BLUE),
        ("Manik", "Backend & Pipeline Engineer", "Model Persistence (joblib), Caching & APIs", ACCENT_TEAL),
        ("Vrunali", "Frontend & UI/UX Developer", "Streamlit Dashboard, Live Analyzer & UX Polish", ACCENT_PURPLE),
        ("Tejas", "ML QA & Integration Engineer", "3-Class Sentiment Modeling, Pytest Suite & QA", ACCENT_AMBER),
    ]

    for i, (name, role, focus, color) in enumerate(members):
        col_left = 1.0 + i * 2.9
        add_card(s1, col_left, 3.8, 2.7, 2.8, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85), left_stripe=color)
        
        box = s1.shapes.add_textbox(Inches(col_left + 0.25), Inches(4.0), Inches(2.35), Inches(2.4))
        tf_m = box.text_frame
        tf_m.word_wrap = True
        
        pm_name = tf_m.paragraphs[0]
        pm_name.text = name
        pm_name.font.size = Pt(18)
        pm_name.font.bold = True
        pm_name.font.color.rgb = RGBColor(255, 255, 255)

        pm_role = tf_m.add_paragraph()
        pm_role.text = role
        pm_role.font.size = Pt(11)
        pm_role.font.bold = True
        pm_role.font.color.rgb = color
        pm_role.space_before = Pt(4)

        pm_focus = tf_m.add_paragraph()
        pm_focus.text = focus
        pm_focus.font.size = Pt(10)
        pm_focus.font.color.rgb = RGBColor(148, 163, 184)
        pm_focus.space_before = Pt(8)

    s1.notes_slide.notes_text_frame.text = (
        "Welcome team! Today we are walking through the core concept, architecture diagrams, "
        "and our 8-day sprint execution roadmap. Every member has a clear single-ownership role."
    )

    # =========================================================================
    # SLIDE 2: The Big Picture & Motivation (Concept Flow)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_LIGHT)
    add_header(s2, "Project Overview: Problem & The Big Picture")

    add_card(s2, 0.8, 1.3, 11.7, 1.3, left_stripe=ACCENT_BLUE)
    b_top = s2.shapes.add_textbox(Inches(1.1), Inches(1.4), Inches(11.2), Inches(1.1))
    tf_top = b_top.text_frame
    tf_top.word_wrap = True
    p = tf_top.paragraphs[0]
    p.text = "THE CORE CHALLENGE & OPPORTUNITY:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p_body = tf_top.add_paragraph()
    p_body.text = "Patients write 215,000+ reviews online detailing drug effectiveness, nausea, and symptom relief. Reading these manually is impossible. Our NLP platform reads these reviews automatically, classifies patient sentiment, and delivers real-time clinical intelligence to doctors, pharma teams, and patients."
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = PRIMARY
    p_body.space_before = Pt(3)

    diagram_steps = [
        ("Step 1: Patient Feedback", "215k+ Online Reviews", "Patients submit reviews describing symptoms, side effects, and ratings (1-10) on Drugs.com.", ACCENT_RED),
        ("Step 2: AI & NLP Processing", "Text -> Intelligence", "Our system cleans text, extracts key medical terms, and runs machine learning classifiers.", ACCENT_BLUE),
        ("Step 3: Actionable Insights", "Clinical Dashboards", "Doctors and researchers see drug rankings, sentiment trends, and side-effect patterns.", ACCENT_TEAL)
    ]

    for i, (title, sub, desc, col) in enumerate(diagram_steps):
        c_left = 0.8 + i * 4.1
        add_card(s2, c_left, 2.9, 3.5, 4.0, left_stripe=col)

        b_st = s2.shapes.add_textbox(Inches(c_left + 0.25), Inches(3.1), Inches(3.05), Inches(3.6))
        tf_st = b_st.text_frame
        tf_st.word_wrap = True

        p_t = tf_st.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_s = tf_st.add_paragraph()
        p_s.text = sub
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = PRIMARY
        p_s.space_before = Pt(4)

        p_d = tf_st.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(10)

        if i < 2:
            add_arrow(s2, c_left + 3.65, 4.65, width=0.3, height=0.3, color=ACCENT_BLUE)

    s2.notes_slide.notes_text_frame.text = (
        "This slide introduces the project: turning 215k+ unstructured text reviews into structured clinical intelligence."
    )

    # =========================================================================
    # SLIDE 3: System Architecture (3-Tier Layered Diagram)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_LIGHT)
    add_header(s3, "System Architecture: 3-Tier Layered Design")

    layers = [
        ("TIER 1: PRESENTATION LAYER (Streamlit Web Dashboard)", [
            "• app.py: Multi-tab interactive UI (Demo Mode Auto-Banner, File Uploader)",
            "• Tab 1: Dataset Analytics & Model Comparison Leaderboard (Matplotlib / Seaborn)",
            "• Tab 2: Live Review Analyzer (Real-time single-review scoring & confidence progress bars)"
        ], ACCENT_BLUE),
        ("TIER 2: PROCESSING & MACHINE LEARNING ENGINE", [
            "• Scikit-Learn Pipeline (ml_pipeline/): TF-IDF Vectorizer, BaseSentimentModel wrapper & utils",
            "• Transformers (hf_sentiment.py): Hugging Face DistilBERT pipeline for deep contextual NLP",
            "• Distributed PySpark Engine (src/ & main.py): PySpark MLlib for big data batch scaling"
        ], ACCENT_TEAL),
        ("TIER 3: DATA & PERSISTENCE STORAGE LAYER", [
            "• Datasets (data/): 1,000-row sample dataset (data/sample/) & raw full CSVs",
            "• Serialized Artifacts (models/): Pre-trained .joblib model files & vectorizers for fast loading (< 2s)",
            "• Structured Logging: streamlit_app.log & pipeline.log with execution latency tracking"
        ], ACCENT_PURPLE)
    ]

    for i, (layer_title, items, col) in enumerate(layers):
        c_top = 1.35 + i * 1.85
        add_card(s3, 0.8, c_top, 11.7, 1.65, left_stripe=col)

        b_ly = s3.shapes.add_textbox(Inches(1.1), Inches(c_top + 0.12), Inches(11.2), Inches(1.4))
        tf_ly = b_ly.text_frame
        tf_ly.word_wrap = True

        p_t = tf_ly.paragraphs[0]
        p_t.text = layer_title
        p_t.font.size = Pt(12.5)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        for item in items:
            p_i = tf_ly.add_paragraph()
            p_i.text = item
            p_i.font.size = Pt(10)
            p_i.font.color.rgb = TEXT_MAIN
            p_i.space_before = Pt(3)

        if i < 2:
            add_down_arrow(s3, 6.5, c_top + 1.63, width=0.25, height=0.22, color=ACCENT_BLUE)

    s3.notes_slide.notes_text_frame.text = (
        "Our system is organized cleanly into 3 tiers: Presentation (Streamlit), Processing (Scikit-Learn/PySpark), "
        "and Storage (Joblib models and CSV data)."
    )

    # =========================================================================
    # SLIDE 4: End-to-End NLP & Sentiment Pipeline (Visual Workflow)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_LIGHT)
    add_header(s4, "NLP Feature Pipeline & 3-Class Sentiment Engine")

    # Left: NLP Pipeline Flow
    add_card(s4, 0.8, 1.4, 5.7, 5.5, left_stripe=ACCENT_BLUE)
    b_nlp = s4.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(5.2), Inches(5.1))
    tf_nlp = b_nlp.text_frame
    tf_nlp.word_wrap = True

    p = tf_nlp.paragraphs[0]
    p.text = "TEXT TRANSFORMATION PIPELINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    nlp_steps = [
        ("1. Raw Patient Narrative:", "Example: \"Worked amazing for my migraines, but caused slight nausea.\""),
        ("2. Cleaning & HTML Unescape:", "BeautifulSoup unescapes &#039; -> ', regex removes special symbols & noise."),
        ("3. TF-IDF Tokenization:", "Extracts 10,000 most significant unigrams & bigrams, removes stopwords."),
        ("4. Feature Matrix:", "Outputs high-dimensional numerical sparse matrix ready for ML models.")
    ]
    for title, desc in nlp_steps:
        p_t = tf_nlp.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = PRIMARY
        p_t.space_before = Pt(8)
        p_d = tf_nlp.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED

    # Right: 3-Class Sentiment Target Mapping
    add_card(s4, 6.8, 1.4, 5.7, 5.5, left_stripe=ACCENT_TEAL)
    b_cls = s4.shapes.add_textbox(Inches(7.1), Inches(1.6), Inches(5.2), Inches(5.1))
    tf_cls = b_cls.text_frame
    tf_cls.word_wrap = True

    p = tf_cls.paragraphs[0]
    p.text = "3-CLASS SENTIMENT TARGET ENGINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    classes = [
        ("🔴 Negative (Class 0): Ratings 1.0 – 3.0", "Ineffective treatment, severe adverse reactions, patient complaints.", ACCENT_RED),
        ("🟡 Neutral (Class 1): Ratings 4.0 – 6.0", "Moderate effectiveness, manageable side effects, mixed outcomes.", ACCENT_AMBER),
        ("🟢 Positive (Class 2): Ratings 7.0 – 10.0", "High satisfaction, quick symptom relief, strong recommendation.", ACCENT_TEAL)
    ]
    for title, desc, col in classes:
        p_t = tf_cls.add_paragraph()
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = col
        p_t.space_before = Pt(10)
        p_d = tf_cls.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED

    p_note = tf_cls.add_paragraph()
    p_note.text = "⭐ Why 3-Class? Binary classification (Positive vs Negative) distorts middle ratings (4-6). 3-Class modeling accurately captures nuanced clinical feedback."
    p_note.font.size = Pt(9.5)
    p_note.font.bold = True
    p_note.font.color.rgb = PRIMARY
    p_note.space_before = Pt(12)

    s4.notes_slide.notes_text_frame.text = (
        "Shows text tokenization to TF-IDF and our 3-class target mapping (Negative 1-3, Neutral 4-6, Positive 7-10)."
    )

    # =========================================================================
    # SLIDE 5: Machine Learning Models Benchmarked
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_LIGHT)
    add_header(s5, "Machine Learning Algorithms Benchmarked")

    models_info = [
        ("Logistic Regression", "Linear Baseline", "Fast, interpretable, L2-regularized linear model. Primary low-latency baseline.", ACCENT_BLUE),
        ("Multinomial Naive Bayes", "Probabilistic NLP", "Fast probabilistic classifier modeling word frequencies across sparse TF-IDF.", ACCENT_TEAL),
        ("Random Forest", "Ensemble of Trees", "100 decision trees evaluating non-linear feature combinations and word interactions.", ACCENT_PURPLE),
        ("Gradient Boosted Trees", "Boosted Trees", "Sequentially trained decision trees with inverse class weighting for imbalance.", ACCENT_AMBER),
        ("Support Vector Machine", "Margin Classifier", "Kernel-based margin classifier with calibrated probabilities for smooth confidence.", ACCENT_RED),
        ("DistilBERT Transformer", "Deep Learning", "66M parameter pre-trained contextual transformer from Hugging Face for nuanced semantics.", PRIMARY)
    ]

    for i, (m_name, m_tag, m_desc, col) in enumerate(models_info):
        row = i // 3
        col_idx = i % 3
        c_left = 0.8 + col_idx * 3.95
        c_top = 1.4 + row * 2.7
        add_card(s5, c_left, c_top, 3.7, 2.5, left_stripe=col)

        b_m = s5.shapes.add_textbox(Inches(c_left + 0.25), Inches(c_top + 0.15), Inches(3.3), Inches(2.2))
        tf_m = b_m.text_frame
        tf_m.word_wrap = True

        p_mn = tf_m.paragraphs[0]
        p_mn.text = m_name
        p_mn.font.size = Pt(14)
        p_mn.font.bold = True
        p_mn.font.color.rgb = col

        p_mt = tf_m.add_paragraph()
        p_mt.text = m_tag
        p_mt.font.size = Pt(10)
        p_mt.font.bold = True
        p_mt.font.color.rgb = PRIMARY
        p_mt.space_before = Pt(2)

        p_md = tf_m.add_paragraph()
        p_md.text = m_desc
        p_md.font.size = Pt(9.5)
        p_md.font.color.rgb = TEXT_MUTED
        p_md.space_before = Pt(6)

    s5.notes_slide.notes_text_frame.text = (
        "We benchmark 5 classical ML algorithms against Hugging Face DistilBERT to balance speed and accuracy."
    )

    # =========================================================================
    # SLIDE 6: 1.5-Week Lean Sprint Deliverables (Problems -> Solutions)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_LIGHT)
    add_header(s6, "1.5-Week Lean Sprint: Problems & Direct Solutions")

    comparisons = [
        ("Problem: App Freezes on Retraining", "Solution: Pre-trained Model Persistence (.joblib)", "Owner: Manik | Save models to disk & load instantly via @st.cache_resource (< 2s).", ACCENT_TEAL),
        ("Problem: Blank App Startup", "Solution: Bundled 1,000-Row Sample Dataset", "Owner: Samar | Curate data/sample/ so the app works instantly out of the box in Demo Mode.", ACCENT_BLUE),
        ("Problem: Binary Distortion (4-6)", "Solution: 3-Class Sentiment Modeling Engine", "Owner: Tejas | Accurately classify Negative (1-3), Neutral (4-6), and Positive (7-10) reviews.", ACCENT_AMBER),
        ("Problem: No Single-Review Testing", "Solution: Live Review Analyzer UI Tab", "Owner: Vrunali | Interactive text box to test custom sentences with instant confidence progress bars.", ACCENT_PURPLE),
        ("Problem: Fragile Test Suite", "Solution: Automated Self-Contained Pytest Suite", "Owner: Tejas | Build synthetic mock fixtures ensuring 100% test pass rate in < 5 seconds.", ACCENT_RED)
    ]

    for i, (prob, sol, detail, col) in enumerate(comparisons):
        c_top = 1.35 + i * 1.15
        add_card(s6, 0.8, c_top, 11.7, 1.05, left_stripe=col)

        b_c = s6.shapes.add_textbox(Inches(1.1), Inches(c_top + 0.08), Inches(11.2), Inches(0.9))
        tf_c = b_c.text_frame
        tf_c.word_wrap = True

        p_p = tf_c.paragraphs[0]
        p_p.text = f"{prob}  ➔  {sol}"
        p_p.font.size = Pt(11.5)
        p_p.font.bold = True
        p_p.font.color.rgb = col

        p_d = tf_c.add_paragraph()
        p_d.text = detail
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_before = Pt(2)

    s6.notes_slide.notes_text_frame.text = (
        "Each of our 5 sprint features directly solves a critical technical debt item from our baseline audit."
    )

    # =========================================================================
    # SLIDE 7: Team Ownership & Responsibility Matrix
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_LIGHT)
    add_header(s7, "Single-Ownership Team Responsibility Matrix")

    team_matrix = [
        ("Samar", "Project Lead & ML Architect", ACCENT_BLUE, [
            "Curate 1,000-row sample dataset (data/sample/)",
            "Code reviews, PR merges to develop & main",
            "Clean legacy hardcoded developer paths",
            "Clinical edge case validation & release tagging"
        ]),
        ("Manik", "Backend & Pipeline Engineer", ACCENT_TEAL, [
            "Model serialization (joblib save/load)",
            "Streamlit @st.cache_resource caching layer",
            "Single-review inference helper function",
            "Input validation & memory profiling"
        ]),
        ("Vrunali", "Frontend & UI/UX Developer", ACCENT_PURPLE, [
            "Multi-tab Streamlit redesign (Analytics vs Live)",
            "Demo mode auto-detection banner UI",
            "Real-time sentiment score cards & confidence bars",
            "Responsive testing, tooltips & UI polish"
        ]),
        ("Tejas", "ML QA & Integration Engineer", ACCENT_AMBER, [
            "3-class sentiment target mapping & metrics",
            "Synthetic mock DataFrame fixtures (conftest.py)",
            "Pytest unit & integration test suites",
            "Boundary exception testing & QA sign-off"
        ])
    ]

    for i, (name, role, col, tasks) in enumerate(team_matrix):
        c_left = 0.8 + i * 2.95
        add_card(s7, c_left, 1.4, 2.8, 5.5, left_stripe=col)

        b_tm = s7.shapes.add_textbox(Inches(c_left + 0.2), Inches(1.55), Inches(2.45), Inches(5.1))
        tf_tm = b_tm.text_frame
        tf_tm.word_wrap = True

        p_name = tf_tm.paragraphs[0]
        p_name.text = name
        p_name.font.size = Pt(17)
        p_name.font.bold = True
        p_name.font.color.rgb = col

        p_r = tf_tm.add_paragraph()
        p_r.text = role
        p_r.font.size = Pt(10)
        p_r.font.bold = True
        p_r.font.color.rgb = PRIMARY
        p_r.space_before = Pt(2)

        p_lbl = tf_tm.add_paragraph()
        p_lbl.text = "CORE DELIVERABLES:"
        p_lbl.font.size = Pt(9)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = TEXT_MUTED
        p_lbl.space_before = Pt(8)

        for task in tasks:
            p_t = tf_tm.add_paragraph()
            p_t.text = f"• {task}"
            p_t.font.size = Pt(8.5)
            p_t.font.color.rgb = TEXT_MAIN
            p_t.space_before = Pt(4)

    s7.notes_slide.notes_text_frame.text = (
        "Every task has exactly one owner to ensure clear accountability and rapid progress."
    )

    # =========================================================================
    # SLIDE 8: 8-Day Roadmap & Git Collaboration Workflow
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, BG_LIGHT)
    add_header(s8, "8-Day Sprint Roadmap & Git Collaboration Flow")

    # Top: 3-Phase Sprint Timeline
    phases = [
        ("Phase 1: Build (Days 1–4)", "Sample CSV, Joblib Persistence, Live UI Wireframe, 3-Class ML & Mock Tests", ACCENT_BLUE),
        ("Phase 2: Integrate (Day 5)", "Merge to develop, wire UI to backend inference & run end-to-end integration test", ACCENT_TEAL),
        ("Phase 3: QA & Release (Days 6–8)", "Clinical edge testing, memory profiling, legacy path cleanup, QA sign-off & main release", ACCENT_PURPLE)
    ]
    for i, (p_title, p_desc, col) in enumerate(phases):
        c_left = 0.8 + i * 4.0
        add_card(s8, c_left, 1.4, 3.7, 2.2, left_stripe=col)
        b_ph = s8.shapes.add_textbox(Inches(c_left + 0.2), Inches(1.5), Inches(3.3), Inches(2.0))
        tf_ph = b_ph.text_frame
        tf_ph.word_wrap = True
        p_t = tf_ph.paragraphs[0]
        p_t.text = p_title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = col
        p_d = tf_ph.add_paragraph()
        p_d.text = p_desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_before = Pt(6)

    # Bottom: Git Flow Diagram
    add_card(s8, 0.8, 3.85, 11.7, 3.1, left_stripe=ACCENT_AMBER)
    b_gf = s8.shapes.add_textbox(Inches(1.1), Inches(4.0), Inches(11.2), Inches(2.8))
    tf_gf = b_gf.text_frame
    tf_gf.word_wrap = True

    p = tf_gf.paragraphs[0]
    p.text = "GIT WORKFLOW & TEAM QUALITY RULES"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    git_steps = [
        "1. Feature Branches: feature/<developer>-<task-name> (Isolated development)",
        "2. Test-First Rule: Run pytest locally; 100% passing tests required before opening PR",
        "3. PR Code Reviews: Samar reviews diffs (< 300 lines) and merges into develop branch",
        "4. Release on Day 8: develop is merged into main and tagged v1.1.0-lean"
    ]
    for gs in git_steps:
        p_g = tf_gf.add_paragraph()
        p_g.text = f"• {gs}"
        p_g.font.size = Pt(10)
        p_g.font.color.rgb = PRIMARY
        p_g.space_before = Pt(4)

    s8.notes_slide.notes_text_frame.text = (
        "Shows our 3-phase execution timeline and our strict Git branching rules to keep code stable."
    )

    # =========================================================================
    # SLIDE 9: Day 1 Kickoff & Action Checklist (Dark Theme)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, BG_DARK)

    tbar9 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    tbar9.fill.solid()
    tbar9.fill.fore_color.rgb = ACCENT_BLUE
    tbar9.line.fill.background()

    tb9 = s9.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.2))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    p0 = tf9.paragraphs[0]
    p0.text = "NEXT STEPS & IMMEDIATE ACTION ITEMS"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p1 = tf9.add_paragraph()
    p1.text = "Sprint Kickoff: Immediate Day 1 Action Plan"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_before = Pt(4)

    actions = [
        ("Samar", "Project Lead", "Initialize feature branches, invite team members & curate sample dataset", ACCENT_BLUE),
        ("Manik", "Backend ML", "Set up virtualenv, audit ml_pipeline/ & design joblib persistence", ACCENT_TEAL),
        ("Vrunali", "Frontend UI", "Launch Streamlit locally, audit UI bottlenecks & wireframe multi-tab layout", ACCENT_PURPLE),
        ("Tejas", "ML & QA", "Audit test suite, build synthetic mock DataFrame fixtures in conftest.py", ACCENT_AMBER),
    ]

    for i, (name, role, act, col) in enumerate(actions):
        c_left = 1.0 + i * 2.9
        add_card(s9, c_left, 2.3, 2.7, 4.3, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85), left_stripe=col)

        box_a = s9.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.5), Inches(2.35), Inches(3.9))
        tf_a = box_a.text_frame
        tf_a.word_wrap = True

        p_name = tf_a.paragraphs[0]
        p_name.text = name
        p_name.font.size = Pt(17)
        p_name.font.bold = True
        p_name.font.color.rgb = RGBColor(255, 255, 255)

        p_r = tf_a.add_paragraph()
        p_r.text = role
        p_r.font.size = Pt(10)
        p_r.font.bold = True
        p_r.font.color.rgb = col
        p_r.space_before = Pt(2)

        p_lb = tf_a.add_paragraph()
        p_lb.text = "IMMEDIATE DAY 1 GOAL:"
        p_lb.font.size = Pt(9)
        p_lb.font.bold = True
        p_lb.font.color.rgb = RGBColor(148, 163, 184)
        p_lb.space_before = Pt(10)

        p_act = tf_a.add_paragraph()
        p_act.text = act
        p_act.font.size = Pt(10.5)
        p_act.font.color.rgb = RGBColor(241, 245, 249)
        p_act.space_before = Pt(4)

    s9.notes_slide.notes_text_frame.text = (
        "Let's kick off Day 1! Each member will pull from main, create their feature branch, "
        "and begin their Day 1 tasks. Let's open the floor for any questions."
    )

    output_path = "Drug_Review_Sentiment_Presentation.pptx"
    try:
        prs.save("Drug_Review_Sentiment_Project_Presentation.pptx")
        output_path = "Drug_Review_Sentiment_Project_Presentation.pptx"
    except PermissionError:
        prs.save("Drug_Review_Sentiment_Presentation.pptx")
        output_path = "Drug_Review_Sentiment_Presentation.pptx"
    print(f"Presentation regenerated successfully with exactly {len(prs.slides)} slides to {output_path}")

if __name__ == "__main__":
    create_presentation()
